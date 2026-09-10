#!/usr/bin/env python3
"""Generate editable HQ field-progress PowerPoint decks (example + blank)."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

NAVY = RGBColor(0x0E, 0x27, 0x44)
INK = RGBColor(0x1C, 0x2A, 0x3A)
MUTED = RGBColor(0x5B, 0x6B, 0x7C)
TEAL = RGBColor(0x1B, 0x7A, 0x6E)
WHITE = RGBColor(0xFF, 0xFC, 0xF7)
PAPER = RGBColor(0xF4, 0xF1, 0xEA)
GOLD = RGBColor(0xB8, 0x86, 0x0B)
GREEN = RGBColor(0x1F, 0x7A, 0x4D)
AMBER = RGBColor(0xB8, 0x6E, 0x13)
LINE = RGBColor(0xD8, 0xD2, 0xC6)
TEAL_SOFT = RGBColor(0xE4, 0xF3, 0xEF)
GREEN_SOFT = RGBColor(0xE3, 0xF3, 0xEA)
AMBER_SOFT = RGBColor(0xF8, 0xED, 0xD6)
WHITE_CARD = RGBColor(0xFF, 0xFF, 0xFF)
FLAG_BG = RGBColor(0xF8, 0xEE, 0xD3)
FLAG_FG = RGBColor(0x6B, 0x4E, 0x00)

W = Inches(13.333)
H = Inches(7.5)
M = Inches(0.45)
OUT_DIR = Path(__file__).resolve().parents[1] / "playbooks"


def set_font(run, size=14, bold=False, color=INK, name="Calibri"):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def textbox(slide, name, left, top, width, height, lines, valign=MSO_ANCHOR.TOP, align=None):
    """Editable text box (always on top of decorative shapes)."""
    box = slide.shapes.add_textbox(left, top, width, height)
    box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.clear()
    for i, item in enumerate(lines):
        if isinstance(item, str):
            txt, size, bold, color = item, 14, False, INK
        else:
            txt = item[0]
            size = item[1] if len(item) > 1 else 14
            bold = item[2] if len(item) > 2 else False
            color = item[3] if len(item) > 3 else INK
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if align:
            p.alignment = align
        p.space_after = Pt(3)
        run = p.add_run()
        run.text = txt
        set_font(run, size=size, bold=bold, color=color)
    return box


def rect(slide, name, left, top, width, height, fill, line=LINE):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    sh.name = name
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
    return sh


def rounded(slide, name, left, top, width, height, fill, line=LINE):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    sh.name = name
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
    return sh


def slide_bg(slide):
    rect(slide, "bg", 0, 0, W, H, PAPER, None)


def header(slide, d):
    textbox(slide, "hdr_left", M, Inches(0.14), Inches(7.0), Inches(0.3), [(d["product"], 11, True, NAVY)])
    textbox(
        slide,
        "hdr_right",
        Inches(7.2),
        Inches(0.14),
        Inches(5.7),
        Inches(0.3),
        [(d["cycle"], 11, True, MUTED)],
        align=PP_ALIGN.RIGHT,
    )


def footer(slide, n):
    rect(slide, "footer_rule", M, Inches(7.05), Inches(12.43), Inches(0.02), LINE, None)
    textbox(
        slide,
        "footer",
        M,
        Inches(7.08),
        Inches(12.43),
        Inches(0.28),
        [
            (
                f"Patterns only - no named HCPs, patient stories, unpublished dossier data  |  {n} / 5",
                10,
                False,
                MUTED,
            )
        ],
    )


EXAMPLE = {
    "product": "BESREMi | HK PV | PM field pack",
    "cycle": "Biweekly | [dates] | Mid | Internal | De-identified",
    "flag": "SAMPLE - replace before sending to HQ",
    "ask_kicker": "THE ASK - FIRST 30 SECONDS",
    "ask_label": "HQ SHOULD DO THIS",
    "ask": "Concentrate the next two weeks on initiation barriers in priority private accounts. Do not treat call volume as the progress metric.",
    "why_now": "Private-sector PSP expiry (end-2026) is already changing close conversations. Formulary packaging needs field themes, not visit counts, while the evidence window is still open.",
    "rec": "Do: report 3 conversion signals mapped to uptake / formulary / compliance. Do not: send trip logs, named HCPs, or CRM dumps.",
    "snap_formulary": "No change to dossier - themes usable, no new unlock",
    "snap_psp": "Base Case + Private Sector (expires end-2026)",
    "snap_compliance": "Green | PV-only lexicon | ET/MF reactive-only",
    "decision_by": "[YYYY-MM-DD]",
    "owner": "PM, with GM for HQ send",
    "channel": "Standing biweekly HQ update",
    "insight": "Initiation, not awareness, is the stall.",
    "rag": [
        ("UPTAKE VS PLAN", "AMBER", AMBER_SOFT, AMBER, "Awareness holds; initiation is slow.", "Definition: private uptake vs agreed HQ plan - not call volume.", "Vs last cycle: unchanged conversion quality."),
        ("FORMULARY TRAJECTORY", "AMBER", AMBER_SOFT, AMBER, "Field themes are usable; dossier not advanced this cycle.", "Definition: strengthened / delayed / no change to HADF packaging.", "Vs last cycle: no new evidence gap; no new unlock."),
        ("COMPLIANCE", "GREEN", GREEN_SOFT, GREEN, "No uncleared claims. ET/MF stayed reactive-only.", "Definition: materials cleared; no open UMAO/HKAPI issues.", "Vs last cycle: clean. Keep raw notes in the vault."),
    ],
    "signals": [
        ("Affordability is understood; initiation is not.", "Observed: Priority accounts repeat PSP logic, then still do not start.", "So what: Uptake risk sits after affordability, not before it.", "Implication: Next cycle = onboarding friction, not more awareness meetings."),
        ("Price vs traditional IFN is the #1 close-blocker.", "Observed: Same objection across accounts: why pay a premium vs off-label IFN.", "So what: Stay inside approved PV evidence and dosing convenience.", "Implication: No IFN-to-IFN superiority claim. Lexicon response only."),
        ("Two PSP paths are creating two close motions.", "Observed: Base Case vs Private Sector (expires end-2026) used as different close paths.", "So what: Expiry is a sequencing issue for HQ, not a slogan.", "Implication: Do not informal-tweak terms. Confirm eligibility before quoting."),
    ],
    "funnel": ["Awareness - holding", "Intent - holding", "Initiation - STALL", "Persistency - too early"],
    "barriers": [
        "1  Initiation / onboarding after PSP explanation  -> Uptake  | New as the pattern",
        "2  Price vs traditional IFN - lexicon-bound answer only  -> Uptake + compliance  | Unchanged",
        "3  Private Sector PSP clock changing urgency unevenly  -> Access / PSP  | Rising",
        "4  [Fourth barrier or delete row]",
    ],
    "moved": "Moved this cycle: Stopped spreading coverage. Next two weeks: initiation checklist on the priority-account list (titles only; no HCP names on this slide).",
    "options": [
        ("NOT RECOMMENDED", "A. Keep activity reports", "Visit counts and trip notes. Easy to write; HQ cannot see the thesis.", WHITE_CARD),
        ("RECOMMENDED", "B. Signal -> implication -> ask", "This pack. Three patterns, RAG vs plan, one decision date.", TEAL_SOFT),
        ("NOT RECOMMENDED", "C. CRM dump", "Looks rigorous, usually unreadable, often too sensitive to forward.", WHITE_CARD),
    ],
    "premortem": "If this fails: someone still scores the team on call volume, so activity theatre returns.",
    "share_risk": "A named KOL or patient-adjacent story leaks into a forwardable deck.",
    "actions": [
        "1  Rewrite last trip report into this 5-slide pack; strip names  |  PM  |  This week",
        "2  Lock RAG definitions with GM (uptake / formulary / compliance)  |  PM + GM  |  Next HQ cycle",
        "3  Compliance pass: no ET/MF promotion, no uncleared claims  |  PM + Legal/Compliance  |  Before send",
    ],
}

BLANK = {
    "product": "[Product] | [Market] | PM field pack",
    "cycle": "Biweekly | [DD Mon - DD Mon YYYY] | [Low/Mid/High] | Internal | De-identified",
    "flag": "TEMPLATE - fill every [bracket] before sending to HQ",
    "ask_kicker": "THE ASK - FIRST 30 SECONDS",
    "ask_label": "HQ SHOULD DO THIS",
    "ask": "[One sentence: what HQ must decide or note. Lead with the action.]",
    "why_now": "[Window, competitor, PSP clock, or formulary cycle that makes this cycle material.]",
    "rec": "Do: [X]. Do not: [Y].",
    "snap_formulary": "[Strengthened / delayed / no change]",
    "snap_psp": "[Base Case / Private Sector / both] + expiry note",
    "snap_compliance": "[G/A/R] | lexicon | reactive-only status",
    "decision_by": "[YYYY-MM-DD]",
    "owner": "[PM / GM / HQ role]",
    "channel": "[Biweekly HQ update / ad-hoc escalation]",
    "insight": "[Insight title: the stall or the unlock - not a chart title]",
    "rag": [
        ("UPTAKE VS PLAN", "[G/A/R]", AMBER_SOFT, AMBER, "[One line vs agreed plan.]", "Definition: [your HQ plan metric]. Not call volume.", "Vs last cycle: [improved / unchanged / worsened]."),
        ("FORMULARY TRAJECTORY", "[G/A/R]", AMBER_SOFT, AMBER, "[Strengthened / delayed / no change.]", "Definition: HADF packaging status this cycle.", "Vs last cycle: [one line]."),
        ("COMPLIANCE", "[G/A/R]", GREEN_SOFT, GREEN, "[Cleared materials? Open issues? Reactive-only held?]", "Definition: UMAO/HKAPI + lexicon.", "Vs last cycle: [one line]."),
    ],
    "signals": [
        ("[Pattern 1 - not a visit count]", "Observed: [What repeated across accounts.]", "So what: [Which HQ KPI it hits.]", "Implication: [What we will do differently.]"),
        ("[Pattern 2]", "Observed: […]", "So what: […]", "Implication: […]"),
        ("[Pattern 3 or delete card]", "Observed: […]", "So what: […]", "Implication: […]"),
    ],
    "funnel": ["Awareness - [status]", "Intent - [status]", "Initiation - [STALL?]", "Persistency - [status]"],
    "barriers": [
        "1  [Barrier]  -> [Uptake / Formulary / Compliance]  | [New / Unchanged / Rising]",
        "2  [Barrier]  -> […]  | […]",
        "3  [Barrier]  -> […]  | […]",
        "4  [Delete if unused]",
    ],
    "moved": "Moved this cycle: [what changed]. Next two weeks: [concrete action].",
    "options": [
        ("OPTION A", "[Option A title]", "[Cost of this path.]", WHITE_CARD),
        ("RECOMMENDED", "[Option B title]", "[Why this is the recommendation.]", TEAL_SOFT),
        ("OPTION C", "[Option C title]", "[Cost of this path.]", WHITE_CARD),
    ],
    "premortem": "[Most likely reason this pack fails or the stall persists.]",
    "share_risk": "[Compliance / privacy fail mode if raw field notes are forwarded.]",
    "actions": [
        "1  [Action]  |  [Role]  |  [Date]",
        "2  [Action]  |  [Role]  |  [Date]",
        "3  [Action]  |  [Role]  |  [Date]",
    ],
}


def build_slide1(slide, d):
    slide_bg(slide)
    header(slide, d)
    rounded(slide, "flag_bg", M, Inches(0.42), Inches(7.2), Inches(0.34), FLAG_BG, None)
    textbox(slide, "flag", M + Inches(0.12), Inches(0.46), Inches(6.9), Inches(0.26), [(d["flag"], 11, True, FLAG_FG)])
    textbox(slide, "kicker", M, Inches(0.82), Inches(12.4), Inches(0.26), [(d["ask_kicker"], 11, True, TEAL)])
    rounded(slide, "ask_bg", M, Inches(1.12), Inches(12.43), Inches(1.55), NAVY, None)
    textbox(slide, "ask_label", M + Inches(0.2), Inches(1.2), Inches(12.0), Inches(0.24), [(d["ask_label"], 11, True, GOLD)])
    textbox(slide, "ask", M + Inches(0.2), Inches(1.48), Inches(12.0), Inches(1.1), [(d["ask"], 22, True, WHITE)])
    rounded(slide, "why_bg", M, Inches(2.85), Inches(6.05), Inches(1.95), WHITE_CARD)
    rounded(slide, "rec_bg", Inches(6.83), Inches(2.85), Inches(6.05), Inches(1.95), WHITE_CARD)
    textbox(slide, "why_now", M + Inches(0.16), Inches(2.95), Inches(5.75), Inches(1.75), [("WHY NOW", 11, True, MUTED), (d["why_now"], 16, False, INK)])
    textbox(slide, "recommendation", Inches(6.99), Inches(2.95), Inches(5.75), Inches(1.75), [("RECOMMENDATION", 11, True, MUTED), (d["rec"], 16, False, INK)])
    for i, (label, val) in enumerate(
        [("FORMULARY THIS CYCLE", d["snap_formulary"]), ("PSP IN PLAY", d["snap_psp"]), ("COMPLIANCE", d["snap_compliance"])]
    ):
        left = M + Inches(i * 4.3)
        rounded(slide, f"snap_bg_{i+1}", left, Inches(4.95), Inches(3.85), Inches(1.05), WHITE_CARD)
        textbox(slide, f"snap_{i+1}", left + Inches(0.14), Inches(5.02), Inches(3.55), Inches(0.9), [(label, 10, True, MUTED), (val, 13, True, NAVY)])
    textbox(slide, "meta", M, Inches(6.15), Inches(12.43), Inches(0.8), [
        (f"DECISION NEEDED BY:  {d['decision_by']}", 11, True, MUTED),
        (f"OWNER:  {d['owner']}", 11, True, MUTED),
        (f"GOES INTO:  {d['channel']}", 11, True, MUTED),
    ])
    footer(slide, 1)


def build_slide2(slide, d):
    slide_bg(slide)
    header(slide, d)
    textbox(slide, "kicker", M, Inches(0.46), Inches(12.4), Inches(0.26), [("FIVE-SECOND SCORECARD", 11, True, TEAL)])
    textbox(slide, "title", M, Inches(0.72), Inches(12.4), Inches(0.65), [(d["insight"], 26, True, NAVY)])
    for i, (name, label, pill_bg, pill_fg, status, defin, delta) in enumerate(d["rag"]):
        left = M + Inches(i * 4.3)
        rounded(slide, f"rag_bg_{i+1}", left, Inches(1.48), Inches(3.85), Inches(5.35), WHITE_CARD)
        rounded(slide, f"rag_pill_{i+1}", left + Inches(2.35), Inches(1.58), Inches(1.2), Inches(0.3), pill_bg, None)
        textbox(slide, f"rag_name_{i+1}", left + Inches(0.16), Inches(1.58), Inches(2.1), Inches(0.28), [(name, 11, True, MUTED)])
        textbox(slide, f"rag_pill_txt_{i+1}", left + Inches(2.42), Inches(1.6), Inches(1.05), Inches(0.26), [(label, 10, True, pill_fg)])
        textbox(slide, f"rag_body_{i+1}", left + Inches(0.16), Inches(1.95), Inches(3.55), Inches(4.7), [
            (status, 18, True, NAVY), (defin, 13, False, MUTED), (delta, 14, True, INK),
        ])
    footer(slide, 2)


def build_slide3(slide, d):
    slide_bg(slide)
    header(slide, d)
    textbox(slide, "kicker", M, Inches(0.46), Inches(12.4), Inches(0.26), [("FIELD SIGNALS - MAX THREE", 11, True, TEAL)])
    textbox(slide, "title", M, Inches(0.72), Inches(12.4), Inches(0.65), [("What repeated in the field - not who we visited", 26, True, NAVY)])
    for i, (title_t, obs, so, impl) in enumerate(d["signals"]):
        left = M + Inches(i * 4.3)
        rounded(slide, f"sig_bg_{i+1}", left, Inches(1.48), Inches(3.85), Inches(5.35), WHITE_CARD)
        textbox(slide, f"signal_{i+1}", left + Inches(0.14), Inches(1.58), Inches(3.55), Inches(5.1), [
            (title_t, 16, True, NAVY), (obs, 13, False, INK), (so, 13, False, INK), (impl, 13, True, TEAL),
        ])
    footer(slide, 3)


def build_slide4(slide, d):
    slide_bg(slide)
    header(slide, d)
    textbox(slide, "kicker", M, Inches(0.46), Inches(12.4), Inches(0.26), [("CONVERSION, NOT COVERAGE", 11, True, TEAL)])
    textbox(slide, "title", M, Inches(0.72), Inches(12.4), Inches(0.65), [("Where field time dies", 26, True, NAVY)])
    rounded(slide, "funnel_bg", M, Inches(1.48), Inches(5.9), Inches(5.35), WHITE_CARD)
    textbox(slide, "funnel_title", M + Inches(0.16), Inches(1.58), Inches(5.5), Inches(0.28), [("ACCOUNT MOTION (QUALITATIVE)", 11, True, MUTED)])
    widths = [Inches(5.0), Inches(4.2), Inches(3.4), Inches(2.7)]
    tops = [Inches(2.05), Inches(2.95), Inches(3.85), Inches(4.75)]
    lefts = [Inches(0.7), Inches(1.1), Inches(1.5), Inches(2.0)]
    for i, (label, w, t, l) in enumerate(zip(d["funnel"], widths, tops, lefts)):
        stall = "STALL" in label.upper()
        rounded(slide, f"funnel_bar_{i+1}", l, t, w, Inches(0.68), NAVY, GOLD if stall else None)
        textbox(slide, f"funnel_txt_{i+1}", l + Inches(0.08), t + Inches(0.12), w - Inches(0.16), Inches(0.5), [(label, 13, True, WHITE)])
    textbox(slide, "funnel_note", M + Inches(0.16), Inches(5.55), Inches(5.5), Inches(0.55), [
        ("Widths are illustrative. Replace with your qualitative read - do not invent counts.", 11, False, MUTED),
    ])
    rounded(slide, "barrier_bg", Inches(6.7), Inches(1.48), Inches(6.18), Inches(5.35), WHITE_CARD)
    lines = [("RANKED BARRIERS (MAX 4)", 11, True, MUTED)] + [(b, 14, False, INK) for b in d["barriers"]] + [(d["moved"], 14, True, NAVY)]
    textbox(slide, "barriers", Inches(6.86), Inches(1.58), Inches(5.85), Inches(5.1), lines)
    footer(slide, 4)


def build_slide5(slide, d):
    slide_bg(slide)
    header(slide, d)
    textbox(slide, "kicker", M, Inches(0.46), Inches(12.4), Inches(0.26), [("DECISION PAGE", 11, True, TEAL)])
    textbox(slide, "title", M, Inches(0.72), Inches(12.4), Inches(0.65), [("Options, risks, and the HQ ask", 26, True, NAVY)])
    for i, (tag, ttl, body, col) in enumerate(d["options"]):
        left = M + Inches(i * 4.3)
        rounded(slide, f"opt_bg_{i+1}", left, Inches(1.42), Inches(3.85), Inches(2.05), col)
        textbox(slide, f"option_{i+1}", left + Inches(0.14), Inches(1.5), Inches(3.55), Inches(1.85), [
            (tag, 10, True, TEAL), (ttl, 15, True, NAVY), (body, 12, False, INK),
        ])
    rounded(slide, "premortem_bg", M, Inches(3.62), Inches(6.05), Inches(1.25), WHITE_CARD)
    rounded(slide, "risk_bg", Inches(6.83), Inches(3.62), Inches(6.05), Inches(1.25), WHITE_CARD)
    textbox(slide, "premortem", M + Inches(0.14), Inches(3.7), Inches(5.75), Inches(1.05), [
        ("PRE-MORTEM", 11, True, MUTED), (d["premortem"], 13, False, INK),
    ])
    textbox(slide, "share_risk", Inches(6.97), Inches(3.7), Inches(5.75), Inches(1.05), [
        ("SHAREABLE-PACK RISK", 11, True, MUTED), (d["share_risk"], 13, False, INK),
    ])
    rounded(slide, "actions_bg", M, Inches(5.05), Inches(12.43), Inches(1.85), WHITE_CARD)
    textbox(slide, "actions", M + Inches(0.14), Inches(5.12), Inches(12.1), Inches(1.7), [
        ("NEXT ACTIONS  |  role + date, not names", 11, True, MUTED),
    ] + [(a, 14, False, INK) for a in d["actions"]])
    footer(slide, 5)


def build_deck(data, filename):
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]
    slides = [prs.slides.add_slide(blank) for _ in range(5)]
    build_slide1(slides[0], data)
    build_slide2(slides[1], data)
    build_slide3(slides[2], data)
    build_slide4(slides[3], data)
    build_slide5(slides[4], data)
    out = OUT_DIR / filename
    prs.save(out)
    return out


def main():
    example = build_deck(EXAMPLE, "hq-field-progress-slides.pptx")
    blank = build_deck(BLANK, "hq-field-progress-slides-blank.pptx")
    print(f"Wrote {example}")
    print(f"Wrote {blank}")


if __name__ == "__main__":
    main()
