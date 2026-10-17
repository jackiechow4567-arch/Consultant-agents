#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build the MGMT5883 consulting report and estimate double-spaced pages."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import ImageFont

OUT = Path("/workspace/course-work/MGMT5883")
FONT_REG = "/usr/share/fonts/truetype/liberation/LiberationSerif-Regular.ttf"
FONT_BOLD = "/usr/share/fonts/truetype/liberation/LiberationSerif-Bold.ttf"

PAGE_W_IN = 6.5
PAGE_H_IN = 9.0
PT = 12
LINE_PT = 24


def wrap_first_indent(text, font, width_px, indent_px):
    words = text.split()
    lines, cur, limit = [], "", width_px - indent_px
    for w in words:
        trial = w if not cur else cur + " " + w
        if font.getlength(trial) <= limit:
            cur = trial
        else:
            if cur:
                lines.append(cur)
            cur = w
            limit = width_px
    if cur:
        lines.append(cur)
    return lines or [""]


def wrap_lines(text, font, width_px):
    words = text.split()
    lines, cur = [], ""
    for w in words:
        trial = w if not cur else cur + " " + w
        if font.getlength(trial) <= width_px:
            cur = trial
        else:
            if cur:
                lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines or [""]


def estimate(blocks):
    width_px = PAGE_W_IN * 72
    font_r = ImageFont.truetype(FONT_REG, PT)
    font_b = ImageFont.truetype(FONT_BOLD, PT)
    font_title = ImageFont.truetype(FONT_BOLD, 14)
    used_pt = 0.0
    section_start = 0.0
    print("--- pagination estimate (Liberation Serif ~= Times, double, 1-inch) ---")
    for kind, text in blocks:
        if kind == "title":
            lines = wrap_lines(text, font_title, width_px)
            h = len(lines) * 18 + 6
        elif kind == "subtitle":
            lines = wrap_lines(text, font_r, width_px)
            h = len(lines) * 16 + 10
        elif kind == "h":
            if section_start:
                print(f"   section length {((used_pt - section_start) / (PAGE_H_IN * 72)):.2f} p")
            section_start = used_pt
            lines = wrap_lines(text, font_b, width_px)
            h = 10 + len(lines) * LINE_PT
        else:
            lines = wrap_first_indent(text, font_r, width_px, 36)
            h = len(lines) * LINE_PT
        used_pt += h
        pages_here = used_pt / (PAGE_H_IN * 72)
        print(f"{kind:8} {len(lines):3} lines  cum {pages_here:.2f}p  {text[:64]}")
    if section_start:
        print(f"   section length {((used_pt - section_start) / (PAGE_H_IN * 72)):.2f} p")
    pages = used_pt / (PAGE_H_IN * 72)
    words = sum(len(t.split()) for k, t in blocks if k == "p")
    print(f"TOTAL {pages:.2f} pages | BODY WORDS {words}")
    # Conservative Word double-spacing (single ~14pt, double ~28pt)
    cons_lines = sum(
        len(wrap_first_indent(t, font_r, width_px, 36)) for k, t in blocks if k == "p"
    )
    cons_pages = (cons_lines * 28 + 8 * 28) / (PAGE_H_IN * 72)
    print(f"CONSERVATIVE Word double ~ {cons_pages:.2f} pages ({cons_lines} body lines)")


def set_run_font(run, name, size, bold=False, italic=False, color=None):
    run.bold = bold
    run.italic = italic
    run.font.name = name
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = r.makeelement(qn("w:rFonts"), {})
        rPr.append(rFonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(a), name)


def add_p(doc, text, *, size=12, bold=False, italic=False, center=False,
          indent=True, space_before=0, space_after=0, double=True):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE if double else WD_LINE_SPACING.SINGLE
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if indent and not center:
        pf.first_line_indent = Inches(0.5)
    run = p.add_run(text)
    set_run_font(run, "Times New Roman", size, bold=bold, italic=italic)
    return p


def main():
    title = "Employee Creativity and Innovation at PharmaEssentia Hong Kong"
    subtitle = "A consulting report for the Hong Kong leadership team  |  October 2026"

    intro = (
        "PharmaEssentia Hong Kong commercializes BESREMi (ropeginterferon alfa-2b) for polycythemia vera "
        "and carries the clinical, medical, and compliance work that accompanies that launch. The subsidiary "
        "is small, specialist, and governed by the Undesirable Medical Advertisements Ordinance and the HKAPI "
        "Code of Practice. In this report, creativity means the generation of new and useful ideas, and innovation "
        "means the successful implementation of those ideas. I identify the obstacles that hinder employees from "
        "doing this work, then recommend one feasible action for each of the two most critical obstacles, with the "
        "steps I would use to put those actions in place. The standard on claims, safety, and conduct stays in "
        "place. The change is a path on which useful local ideas can be voiced, tested, and decided."
    )

    obs_open = (
        "Other conditions also hinder creativity: little slack, narrow associate roles, and little joint work between "
        "commercial and clinical colleagues. In Amabile's work-environment model these are workload pressure and "
        "evaluation pressure. Two obstacles matter more."
    )

    obs1 = (
        "The first is a split between local ideas and the authority to implement them. Staff can propose useful changes "
        "to patient support, formulary evidence, investigator engagement, and handoffs. Claims, price corridors, contracts, "
        "and portfolio choices are decided at headquarters, and Hong Kong cannot sign every commitment. A local explanation "
        "waits for global lexicon clearance before the team may use it, so the person who sees the gap does not carry the "
        "idea into implementation. Creativity and innovation come apart. Administrative innovation follows when leaders "
        "design both creation and implementation (Chen, Gong, Song, and Wang, 2021). The decision map is the evidence: "
        "Hong Kong owns local tactics; headquarters owns the lexicon, price corridors, and contracts; and no routine turns "
        "one local idea into one headquarters decision."
    )

    obs2 = (
        "The second obstacle is more critical because it stops ideas earlier. Success means delivery against plan, progress "
        "on public access, and a clean promotional and safety record. That is the right floor for a regulated medicine. It "
        "hinders creativity when it is the only signal people receive: a novel comment is heard as a compliance risk, and "
        "goals reward the number and the code. Learning orientation and psychological safety raise creativity by making "
        "information exchange possible (Gong, Kim, Zhu, and Lee, 2013; Gong, Cheung, Wang, and Huang, 2012). Creative "
        "self-efficacy and intrinsic motivation for creativity strengthen each other. On difficult work, encouragement and "
        "small wins build the belief that a creative result is possible (Chen, Kim, Gong, and Liang, 2025). Feedback here "
        "is clearance, lexicon checks, and target reviews, so staff conclude that accurate execution is what the organization values."
    )

    rec1 = (
        "For the implementation gap, I recommend a quarterly idea-to-decision path. Decision rights stay with "
        "headquarters. Once a quarter, the Head of Office chairs a 90-minute session with product, clinical, and "
        "compliance in the same room. The technique is to defer judgment, go for quantity, and build on other "
        "people's ideas. Each idea is one page: the local problem, why the idea is new, why it is useful for patients "
        "or for Hong Kong, and which claims it leaves untouched. Price, new scientific claims, and contract language "
        "stay off that page. Compliance attends in order to design a safe test. The group keeps the single strongest "
        "idea and scores it on usefulness, compliance containment, and feasibility. The Head of Office then sends "
        "headquarters the note this office already uses: two or three options, one recommendation, and a clear split "
        "between what Hong Kong can do alone and what needs medical, regulatory, or commercial approval."
    )

    rec1b = (
        "I would apply the course emphasis on disruptive innovation as a boundary (Christensen, Raynor, and McDonald, "
        "2015). The scientific core of BESREMi is a sustaining innovation owned by headquarters, and the processes "
        "that protect claims and contracts exist to keep that core consistent across markets. Local experiments belong "
        "in the operating system around the product: how evidence is prepared for formulary dialogue, how a field "
        "question reaches medical information, and how patient-support steps reduce drop-off. In week one, the Head "
        "of Office states that an idea remains an idea until the second stage classifies it. In week two, the pilot "
        "uses one problem that commercial and clinical colleagues already share. Within two further weeks, the note "
        "goes to headquarters. At day 90, I would apply one test. The idea has received a decision: a local pilot, a "
        "request to revise, or a decline that states the reason. A reasoned decline teaches the system. An idea that "
        "disappears means the path has failed."
    )

    rec2 = (
        "For the climate, I recommend one covered creative assignment for every associate, each quarter. Performance "
        "goals and compliance goals remain, and a learning goal is added beside them. Intrinsic interest and extrinsic "
        "goals produce more creativity when they are tuned together, so the assignment is real work and the attempt "
        "is recognized (Gong, Wu, Song, and Zhang, 2017). A commercial colleague might redraw the path a field question "
        "takes into medical information. A clinical colleague might propose two cleaner handoffs between site follow-up "
        "and medical information. The manager tells the person that they are capable of the work, defines the finished "
        "product as a one-page set of options, and comments on novelty and usefulness. The task sits slightly outside "
        "the person's usual lane. That variety engages them in the creative process and rebuilds creative self-efficacy "
        "on difficult work (Chen, Kim, Gong, and Liang, 2025). The first review uses the same rules as the quarterly "
        "session, so comments address the idea in the convergent pass. One line is added to the biweekly review the "
        "office already holds: what we learned, and what we will try next."
    )

    rec2b = (
        "I would implement this with the resources the office already has. Headcount stays as it is, and compliance "
        "remains above commercial convenience. The lever is managerial behavior on hard work: encouragement, a bounded "
        "challenge, and a small win. By the end of the quarter, each associate has completed one small win, and the "
        "Head of Office can point to one process that changed because of an employee's idea. A five-question pulse at "
        "that point, covering voice, learning, safety, time, and whether an idea received a decision, tells leadership "
        "whether the climate moved."
    )

    close = (
        "PharmaEssentia Hong Kong has a serious task and capable people. The subsidiary needs a path from a new and "
        "useful idea to a decision, and a climate in which employees can develop an idea and then have that idea judged. "
        "The two actions are small on purpose. A quarterly idea-to-decision path makes innovation mean implementation. "
        "One covered assignment per person makes creative work part of the job. Claims, safety, and conduct remain the "
        "floor. Headquarters receives an option and a recommendation. The people closest to patients in Hong Kong gain "
        "a way to improve the system around a medicine that already has a global standard."
    )

    blocks = [
        ("title", title),
        ("subtitle", subtitle),
        ("h", "Introduction"),
        ("p", intro),
        ("h", "Obstacles to creativity and innovation"),
        ("p", obs_open),
        ("p", obs1),
        ("p", obs2),
        ("h", "Recommendations and how I would implement them"),
        ("p", rec1),
        ("p", rec1b),
        ("p", rec2),
        ("p", rec2b),
        ("h", "Conclusion"),
        ("p", close),
    ]
    estimate(blocks)
    for name, text in (
        ("intro", intro),
        ("obstacles", " ".join([obs_open, obs1, obs2])),
        ("recs", " ".join([rec1, rec1b, rec2, rec2b])),
        ("close", close),
    ):
        print(f"WORDS {name:10} {len(text.split())}")

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)

    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("PharmaEssentia Hong Kong  |  Consulting report")
    set_run_font(r, "Times New Roman", 9, italic=True, color=RGBColor(0x55, 0x55, 0x55))
    header.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("Course draft  |  Page ")
    set_run_font(r, "Times New Roman", 9, color=RGBColor(0x55, 0x55, 0x55))
    fld = footer.add_run()
    fldChar1 = fld._r.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): "begin"})
    instr = fld._r.makeelement(qn("w:instrText"), {})
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fldChar2 = fld._r.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): "end"})
    fld._r.append(fldChar1)
    fld._r.append(instr)
    fld._r.append(fldChar2)
    set_run_font(fld, "Times New Roman", 9, color=RGBColor(0x55, 0x55, 0x55))

    add_p(doc, title, size=14, bold=True, center=True, indent=False, double=False, space_after=2)
    add_p(doc, subtitle, size=11, italic=True, center=True, indent=False, double=False, space_after=6)

    sections = [
        ("Introduction", [intro]),
        ("Obstacles to creativity and innovation", [obs_open, obs1, obs2]),
        ("Recommendations and how I would implement them", [rec1, rec1b, rec2, rec2b]),
        ("Conclusion", [close]),
    ]
    for h, paras in sections:
        add_p(doc, h, size=12, bold=True, indent=False, space_before=8, double=True)
        for para in paras:
            add_p(doc, para, size=12, indent=True, double=True)

    docx_path = OUT / "consulting-report.docx"
    doc.save(docx_path)
    print("WROTE", docx_path)

    md_parts = [f"# {title}", "", f"*{subtitle}*", ""]
    for h, paras in sections:
        md_parts.append(f"## {h}")
        md_parts.append("")
        for para in paras:
            md_parts.append(para)
            md_parts.append("")
    md_path = OUT / "consulting-report.md"
    md_path.write_text("\n".join(md_parts), encoding="utf-8")
    print("WROTE", md_path)

    html = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>{title}</title>
<style>
  @page {{ size: letter; margin: 1in; }}
  body {{
    font-family: "Times New Roman", Times, serif;
    font-size: 12pt;
    line-height: 2;
    max-width: 6.5in;
    margin: 1in auto;
    color: #111;
  }}
  h1 {{ font-size: 14pt; text-align: center; line-height: 1.3; margin: 0 0 0.2em; }}
  .sub {{ text-align: center; font-style: italic; font-size: 11pt; line-height: 1.3; margin-bottom: 0.6em; }}
  h2 {{ font-size: 12pt; line-height: 2; margin: 0.4em 0 0; }}
  p {{ margin: 0; text-indent: 0.5in; }}
</style>
</head>
<body>
<h1>{title}</h1>
<p class="sub">{subtitle}</p>
"""
    for h, paras in sections:
        html += f"<h2>{h}</h2>\n"
        for para in paras:
            html += f"<p>{para}</p>\n"
    html += "</body></html>\n"
    html_path = OUT / "consulting-report.html"
    html_path.write_text(html, encoding="utf-8")
    print("WROTE", html_path)


if __name__ == "__main__":
    main()
